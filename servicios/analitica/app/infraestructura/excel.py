"""Libro de Excel de un reporte (spec 015, paso 9): hojas "Datos" y "Parámetros"."""

from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.configuracion import ZONA

FORMATOS = {
    "MONTO": "#,##0.00",
    "PORCENTAJE": "0.0%",
    "FECHA": "dd/mm/yyyy",
    "INSTANTE": "dd/mm/yyyy hh:mm",
    "NUMERO": "#,##0.##",
}


# Inyección de fórmulas (CSV/Excel injection): `razonSocial` y otros textos los escribe cualquier cliente.
PELIGROSOS = ("=", "+", "-", "@", "\t", "\r")


def _texto_seguro(hoja: Any) -> None:
    """Toda celda de texto que podría leerse como fórmula queda como texto (`data_type = s` y `quotePrefix`)."""
    for fila in hoja.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str) and celda.value.startswith(PELIGROSOS):
                celda.data_type = "s"
                celda.quotePrefix = True


def _celda(valor: Any, tipo: str) -> Any:
    if valor is None:
        return None
    if tipo == "MONTO":
        return Decimal(str(valor))
    if tipo == "FECHA":
        return date.fromisoformat(valor) if isinstance(valor, str) else valor
    if tipo == "INSTANTE":
        instante = datetime.fromisoformat(valor.replace("Z", "+00:00")) if isinstance(valor, str) else valor
        return instante.astimezone(ZONA).replace(tzinfo=None)
    if tipo == "BOOLEANO":
        return "Sí" if valor else "No"
    return valor


def libro(resultado: dict[str, Any], sub: str, ahora: datetime) -> bytes:
    wb = Workbook()
    datos = wb.active
    datos.title = "Datos"
    columnas = resultado["columnas"]
    negrita = Font(bold=True)
    fondo = PatternFill("solid", fgColor="DDE7F3")
    datos.append([c["titulo"] for c in columnas])
    for celda in datos[1]:
        celda.font, celda.fill = negrita, fondo
    for fila in resultado["filas"]:
        datos.append([_celda(fila.get(c["clave"]), c["tipo"]) for c in columnas])
    totales = resultado.get("totales") or {}
    if any(c["clave"] in totales for c in columnas):
        datos.append([_celda(totales.get(c["clave"]), c["tipo"]) if c["clave"] in totales else None for c in columnas])
        datos.cell(datos.max_row, 1, "Total")
        for celda in datos[datos.max_row]:
            celda.font = negrita
    for i, c in enumerate(columnas, start=1):
        letra = get_column_letter(i)
        if c["tipo"] in FORMATOS:
            for (celda,) in datos.iter_rows(min_row=2, min_col=i, max_col=i):
                celda.number_format = FORMATOS[c["tipo"]]
        datos.column_dimensions[letra].width = max(12, len(c["titulo"]) + 4)
    datos.freeze_panes = "A2"

    parametros = wb.create_sheet("Parámetros")
    datos_al = resultado.get("datosAl")
    for clave, valor in [
        ("Reporte", resultado["titulo"]),
        *((f"Parámetro {k}", str(v)) for k, v in resultado["parametros"].items()),
        ("Generado en", ahora.astimezone(ZONA).replace(tzinfo=None)),
        ("Generado por", sub),
        ("Datos al", _celda(datos_al, "INSTANTE") if datos_al else "Sin eventos procesados"),
    ]:
        parametros.append([clave, valor])
    for (celda,) in parametros.iter_rows(min_col=1, max_col=1):
        celda.font = negrita
    for fila in parametros.iter_rows(min_col=2, max_col=2):
        if isinstance(fila[0].value, datetime):
            fila[0].number_format = FORMATOS["INSTANTE"]
    parametros.column_dimensions["A"].width = 24
    parametros.column_dimensions["B"].width = 48
    for hoja in (datos, parametros):
        _texto_seguro(hoja)

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def ahora() -> datetime:
    return datetime.now(UTC)
